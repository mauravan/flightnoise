"""Importiert die EASA-ANP-Datenbank (neun .xlsx-Dateien) nach data/anp.sqlite.

Verwendung (vom Projekt-Root aus)::

    .venv/bin/python data/import_anp.py

Schema: eine Tabelle pro Workbook, Spaltennamen aus der Kopfzeile (Zeile 1)
abgeleitet (kleingeschrieben, nicht-alphanumerisch zu ``_``). Spalten sind
untypisiert (SQLite dynamic typing; int/float/str/None werden verbatim
gespeichert). Tabellen und Primärschlüssel:

- aircraft (acft_id)
- npd_data (npd_id, noise_metric, op_mode, power_setting) + Index
  idx_npd_mode (npd_id, op_mode, noise_metric)
- aerodynamic_coefficients (acft_id, op_type, flap_id)
- default_approach_procedural_steps (acft_id, profile_id, step_number, flap_id)
- default_departure_procedural_steps (acft_id, profile_id, stage_length,
  step_number)
- default_fixed_point_profiles (acft_id, op_type, profile_id, stage_length,
  point_number)
- default_weights (acft_id, stage_length)
- jet_engine_coefficients (acft_id, thrust_rating)
- spectral_classes (spectral_class_id)

Der Import ist idempotent/deterministisch (DROP TABLE IF EXISTS pro Tabelle).
Sowohl die Rohdaten (``data/easa_anp/`` .xlsx/.zip) als auch die erzeugte
``data/anp.sqlite`` sind gitignoriert (EASA-Nutzungsbedingungen verbieten
Weiterverbreitung); nur dieses Skript und das Schema sind committed.
"""

from __future__ import annotations

import itertools
import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

SRC_DIR = Path("data/easa_anp")
DB_PATH = Path("data/anp.sqlite")

# (filename, table_name, primary-key columns); je Workbook genau ein Sheet.
SOURCES: list[tuple[str, str, list[str]]] = [
    ("EASA_ANP_database_Aircraft_v9.xlsx", "aircraft", ["acft_id"]),
    (
        "EASA_ANP_database_NPD_Data_v9.xlsx",
        "npd_data",
        ["npd_id", "noise_metric", "op_mode", "power_setting"],
    ),
    (
        "EASA_ANP_database_Aerodynamic_Coefficients_v9.xlsx",
        "aerodynamic_coefficients",
        ["acft_id", "op_type", "flap_id"],
    ),
    (
        "EASA_ANP_database_Default_Approach_Procedural_Steps_v9.xlsx",
        "default_approach_procedural_steps",
        ["acft_id", "profile_id", "step_number", "flap_id"],
    ),
    (
        "EASA_ANP_database_Default_Departure_Procedural_Steps_v9.xlsx",
        "default_departure_procedural_steps",
        ["acft_id", "profile_id", "stage_length", "step_number"],
    ),
    (
        "EASA_ANP_database_Default_Fixed_Point_Profiles_v9.xlsx",
        "default_fixed_point_profiles",
        ["acft_id", "op_type", "profile_id", "stage_length", "point_number"],
    ),
    (
        "EASA_ANP_database_Default_Weights_v9.xlsx",
        "default_weights",
        ["acft_id", "stage_length"],
    ),
    (
        "EASA_ANP_database_Jet_Engine_Coefficients_v9.xlsx",
        "jet_engine_coefficients",
        ["acft_id", "thrust_rating"],
    ),
    (
        "EASA_ANP_database_Spectral_Classes_v9.xlsx",
        "spectral_classes",
        ["spectral_class_id"],
    ),
]


def sanitize(name: str) -> str:
    """Kopfzelltext -> Spaltenname (lower, nicht-alphanumerisch zu _)."""
    return re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_")


def clean(value: object) -> object:
    """None -> NULL; str gestrippt (leer -> NULL); Zahlen unverändert."""
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def read_sheet(path: Path) -> tuple[list[str], list[tuple]]:
    """Liest das erste Sheet: (Spaltennamen, bereinigte Datenzeilen)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        # Leere formatierte Zellen am Zeilenende abschneiden.
        ncols = len(header)
        while ncols > 0 and header[ncols - 1] is None:
            ncols -= 1
        columns = [sanitize(str(h)) for h in header[:ncols]]
        rows = [
            tuple(clean(v) for v in itertools.islice(row, ncols))
            for row in ws.iter_rows(min_row=2, values_only=True)
        ]
    finally:
        wb.close()
    return columns, rows


def import_anp(src_dir: Path = SRC_DIR, db_path: Path = DB_PATH) -> dict[str, int]:
    """Liest alle neun Workbooks und baut die SQLite-DB neu auf."""
    counts: dict[str, int] = {}
    conn = sqlite3.connect(db_path)
    try:
        for filename, table, pk in SOURCES:
            columns, rows = read_sheet(src_dir / filename)
            conn.execute(f"DROP TABLE IF EXISTS {table}")
            pk_clause = f", PRIMARY KEY ({', '.join(pk)})"
            conn.execute(f"CREATE TABLE {table} ({', '.join(columns)}{pk_clause})")
            conn.executemany(
                f"INSERT INTO {table} VALUES ({', '.join('?' * len(columns))})",
                rows,
            )
            conn.commit()
            counts[table] = len(rows)
        conn.execute(
            "CREATE INDEX idx_npd_mode ON npd_data (npd_id, op_mode, noise_metric)"
        )
        conn.commit()
    finally:
        conn.close()
    return counts


def main() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    counts = import_anp()
    total = 0
    for _filename, table, _pk in SOURCES:
        n = counts[table]
        total += n
        print(f"{table}: {n} rows")
    print(f"total: {total} rows")


if __name__ == "__main__":
    main()
