"""Importiert die EASA-Ersatztabelle (ANP Aircraft Substitutions – Jets & heavy
props, Liste 22-02-2018) in die Tabelle ``substitutions`` von data/anp.sqlite.

Verwendung (vom Projekt-Root aus)::

    .venv/bin/python data/import_substitutions.py

Gelesen wird das Sheet ``by ICAO code`` (Spalten per Kopftext, nie per
Position): ``ICAO_CODE`` → ``ANP_PROXY`` plus die vorgeschriebenen Zuschläge
``DELTA_DEP_dB``/``DELTA_APP_dB`` (ICAO Doc 9911 / ECAC Doc 29). Das Sheet
``by aircraft configuration`` wird nicht importiert — es braucht eine
Motorvariante, die der Verkehrsstichprobe (tar1090-Typecode) fehlt.

Nur Flugzeuge, die in ANP v9 oder im ANP-Archiv v2.3 existieren, sind als
Proxy brauchbar; die Filterung übernimmt ``AnpDatabase.substitutions`` zur
Abfragezeit. ``AIRCRAFT_VARIANT``, ``N_DEP``/``N_APP`` (äquivalente Bewegungen)
werden nicht übernommen. xlsx und DB sind gitignoriert
(EASA-Nutzungsbedingungen verbieten Weiterverbreitung).
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import load_workbook

TABLE = "substitutions"
DDL = ("CREATE TABLE substitutions ("
       "icao_code TEXT, acft_id TEXT, delta_dep_db REAL, delta_app_db REAL, "
       "PRIMARY KEY (icao_code, acft_id))")
SHEET = "by ICAO code"
_SUBSTITUTIONS_XLSX_NAME = (
    "data/easa_anp/anp_aircraft_substitutions_-_jets_heavy_props_22022018_.xlsx"
)
SUBSTITUTIONS_URL = "https://www.easa.europa.eu/en/downloads/138165/en"
XLSX_PATH = Path(_SUBSTITUTIONS_XLSX_NAME)
DB_PATH = Path("data/anp.sqlite")


def import_substitutions(
    xlsx: Path = XLSX_PATH, db_path: Path = DB_PATH
) -> int:
    """Liest Sheet 'by ICAO code'; gibt die Anzahl geschriebener Zeilen zurück."""
    if not xlsx.exists():
        raise FileNotFoundError(
            f"{xlsx} nicht gefunden; 'fluglaerm setup' holt die Datei "
            "(curl -L -o " + _SUBSTITUTIONS_XLSX_NAME + " " + SUBSTITUTIONS_URL
            + "), dann data/import_substitutions.py erneut ausführen."
        )
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise ValueError(
                f"{xlsx}: Sheet '{SHEET}' fehlt (gefunden: {wb.sheetnames})"
            )
        ws = wb[SHEET]
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        try:
            i_icao = header.index("ICAO_CODE")
            i_proxy = header.index("ANP_PROXY")
            i_dep = header.index("DELTA_DEP_dB")
            i_app = header.index("DELTA_APP_dB")
        except ValueError as exc:
            raise ValueError(
                f"{xlsx} ('{SHEET}'): Spalte fehlt in der Kopfzeile: {exc}"
            ) from exc
        records: list[tuple[str, str, float, float]] = []
        for row in rows_iter:

            def cell(idx: int) -> object:
                return row[idx] if idx < len(row) else None

            icao = str(cell(i_icao) or "").strip()
            proxy = str(cell(i_proxy) or "").strip()
            if not icao or not proxy:
                continue
            records.append(
                (
                    icao,
                    proxy,
                    float(cell(i_dep) or 0.0),
                    float(cell(i_app) or 0.0),
                )
            )
    finally:
        wb.close()
    records.sort(key=lambda r: (r[0], r[1]))
    conn = sqlite3.connect(db_path)
    try:
        conn.execute(f"DROP TABLE IF EXISTS {TABLE}")
        conn.execute(DDL)
        conn.executemany(
            f"INSERT OR REPLACE INTO {TABLE} VALUES (?, ?, ?, ?)", records
        )
        conn.commit()
    finally:
        conn.close()
    return len(records)


def main() -> None:
    n = import_substitutions()
    print(f"{TABLE}: {n} rows")


if __name__ == "__main__":
    main()
